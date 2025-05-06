import os
import pandas as pd
from openpyxl import load_workbook

def process_excel_files(directory):
    # 分类规则映射字典
    category_map = {
        "法律": ("一、法律", "1.1 法律"),
        "司法解释": ("一、法律", "1.2 司法解释"),
        "立法解释": ("一、法律", "1.3 立法解释"),
        "行政法规": ("二、规章", "2.1 国务院 - 规章"),
        "行政法规解释": ("二、规章", "2.1 国务院 - 规章"),
        "部委规章": ("二、规章", "2.2 国务院部委办局 - 规章"),
        "国务院部门规章库 - 部委规章": ("二、规章", "2.2 国务院部委办局 - 规章"),
        "地方规章": ("二、规章", "2.3 省市县政府 - 规章"),
        "地方政府规章": ("二、规章", "2.3 省市县政府 - 规章"),
        "地方法规": ("二、规章", "2.3 省市县政府 - 规章"),
        "自治条例、单行条例": ("二、规章", "2.3 省市县政府 - 规章"),
        "军事法规": ("二、规章", "2.4 军事 - 规章"),
        "规范性文件": ("五、规范", "5.1 国务院 - 规范文件"),
        "国务院规范性文件": ("五、规范", "5.1 国务院 - 规范文件"),
        "国务院部门规章库 - 规范性文件": ("五、规范", "5.2 国务院部委办局 - 规范文件"),
        "部门规范性文件": ("五、规范", "5.2 国务院部委办局 - 规范文件"),
        "地方规范性文件": ("五、规范", "5.3 省市县政府 - 规范文件"),
        "地方司法规范": ("五、规范", "5.3 省市县政府 - 规范文件"),
        "团体/行业规范": ("五、规范", "5.4 其他 - 规范文件"),
        "国际法、国际条约与惯例及工作文件": ("六、文件", "6.1 国务院 - 工作文件"),
        "国务院工作文件": ("六、文件", "6.1 国务院 - 工作文件"),
        "工作文件": ("六、文件", "6.1 国务院 - 工作文件"),
        "部门工作文件": ("六、文件", "6.2 国务院部委办局 - 工作文件"),
        "地方工作文件": ("六、文件", "6.3 省市县政府 - 工作文件"),
        "司法指导性文件": ("六、文件", "6.4 司法指导性文件"),
        "其他": ("七、其他", "其他")
    }

    for filename in os.listdir(directory):
        if filename.endswith('.xlsx'):
            filepath = os.path.join(directory, filename)
            
            try:
                # 使用openpyxl加载工作簿
                wb = load_workbook(filepath)
                ws = wb.active
                
                # 获取U列第二行的值（索引从1开始）
                u_cell_value = ws['U2'].value if ws['U2'] else None
                
                if u_cell_value in category_map:
                    x_header = "大类"
                    y_header = "二级分类"
                    
                    # 添加表头
                    ws['X1'] = x_header
                    ws['Y1'] = y_header
                    
                    # 填充分类数据
                    x_value, y_value = category_map[u_cell_value]
                    for row in ws.iter_rows(min_row=2):
                        ws.cell(row=row[0].row, column=24).value = x_value  # X列
                        ws.cell(row=row[0].row, column=25).value = y_value  # Y列
                
                # 保存修改
                wb.save(filepath)
                print(f"成功处理文件：{filename}")
                
            except Exception as e:
                print(f"处理文件 {filename} 时出错：{str(e)}")

def process_all_folders(root_dir):
    """递归处理所有子文件夹中的Excel文件"""
    for foldername, subfolders, filenames in os.walk(root_dir):
        print(f"\n正在处理文件夹: {foldername}")
        # 过滤出Excel文件
        excel_files = [f for f in filenames if f.endswith('.xlsx')]
        if excel_files:
            process_excel_files(foldername)

# 方法一：使用绝对路径（推荐）
target_directory = os.path.join(os.getcwd(), "部委规章-cp")

# 方法二：相对路径写法
target_directory = "./部委规章-cp"  # 或 "部委规章-cp"

# 在现有代码中修改
if __name__ == "__main__":
    # 使用绝对路径写法（示例）
    target_directory = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),  # 获取脚本所在目录
        "0506四轮最后结果"
    )
    process_all_folders(target_directory)  # 改为调用新的入口函数